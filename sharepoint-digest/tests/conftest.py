"""Shared test helpers: sample Office files, a fake Microsoft Graph, and a fake chat model."""

from __future__ import annotations

import io
import json
import os
import threading
import time
from datetime import datetime
from pathlib import Path

from azure.core.credentials import AccessToken
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from sharepoint_digest.config import FoundrySettings, Settings, SharePointSettings
from sharepoint_digest.sharepoint import GRAPH_URL, GraphClient

SITE_URL = "https://contoso.sharepoint.com/sites/Team"


def make_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title and content
    slide.shapes.title.text = "Q3 plan"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue up 12%"
    detail = body.add_paragraph()
    detail.text = "EMEA drove most of the growth"
    detail.level = 1
    slide.notes_slide.notes_text_frame.text = "Mention the hiring freeze"

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Numbers"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (row, col), text in {(0, 0): "Region", (0, 1): "Sales", (1, 0): "EMEA", (1, 1): "4.2M"}.items():
        table.cell(row, col).text = text
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("2026", (1.5, 2))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue by quarter"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(1), Inches(2), Inches(1)).text_frame.text = "Grouped callout"

    prs.slides.add_slide(prs.slide_layouts[6])  # blank: should be left out
    return _save(prs)


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Project Alpha status", 0)
    doc.add_heading("Summary", 1)
    doc.add_paragraph("We are on track for the October launch.")
    doc.add_paragraph("Finalize vendor contract", style="List Bullet")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 2).text = "Date"
    table.cell(1, 0).text = "Beta"
    table.cell(1, 1).text = "Internal"
    table.cell(1, 2).text = "2026-10-01"
    doc.add_paragraph("Next review in two weeks.")
    return _save(doc)


def write_file(root: Path, relative: str, content: bytes, modified: str) -> Path:
    """Create root/relative with the given content and modified time (ISO 8601)."""
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    stamp = datetime.fromisoformat(modified).timestamp()
    os.utime(path, (stamp, stamp))
    return path


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeResponse:
    def __init__(self, url: str, status_code: int = 200, payload=None, content: bytes = b""):
        self.url = url
        self.status_code = status_code
        self.reason = "OK" if status_code < 400 else "Error"
        self._payload = payload
        self.content = json.dumps(payload).encode() if payload is not None else content

    @property
    def text(self) -> str:
        return self.content.decode("utf-8", "replace")

    def json(self):
        if self._payload is None:
            raise ValueError("not JSON")
        return self._payload


class FakeSession:
    """Serves canned Graph responses by URL. Query parameters are recorded, not matched."""

    def __init__(self, routes: dict):
        self.routes = routes
        self.calls: list[dict] = []

    def get(self, url, headers=None, params=None, timeout=None):
        self.calls.append({"url": url, "params": params, "headers": headers})
        route = self.routes.get(url)
        if route is None:
            error = {"error": {"code": "itemNotFound", "message": "The resource could not be found."}}
            return FakeResponse(url, 404, error)
        if isinstance(route, FakeResponse):
            return route
        if isinstance(route, bytes):
            return FakeResponse(url, content=route)
        return FakeResponse(url, payload=route)


class FakeCredential:
    def __init__(self):
        self.tokens_issued = 0

    def get_token(self, *scopes, **kwargs):
        self.tokens_issued += 1
        return AccessToken(f"token-{self.tokens_issued}", int(time.time()) + 3600)


def fake_graph(routes: dict) -> tuple[GraphClient, FakeSession, FakeCredential]:
    session = FakeSession(routes)
    credential = FakeCredential()
    return GraphClient(credential, session_factory=lambda: session), session, credential


def site_routes() -> dict:
    """Routes for a site with two libraries and a "Temp Folder 1" folder in "Documents"."""
    return {
        f"{GRAPH_URL}/sites/contoso.sharepoint.com:/sites/Team": {"id": "site-1"},
        f"{GRAPH_URL}/sites/site-1/drives": {
            "value": [
                {"id": "drive-docs", "name": "Documents", "webUrl": f"{SITE_URL}/Shared%20Documents"},
                {"id": "drive-archive", "name": "Archive", "webUrl": f"{SITE_URL}/Archive"},
            ]
        },
        f"{GRAPH_URL}/drives/drive-docs/root:/Temp%20Folder%201": {
            "id": "folder-1",
            "name": "Temp Folder 1",
            "folder": {"childCount": 3},
            "webUrl": f"{SITE_URL}/Shared%20Documents/Temp%20Folder%201",
        },
    }


def children_url(item_id: str) -> str:
    return f"{GRAPH_URL}/drives/drive-docs/items/{item_id}/children"


def content_url(item_id: str) -> str:
    return f"{GRAPH_URL}/drives/drive-docs/items/{item_id}/content"


def graph_file(item_id: str, name: str, modified: str, *, created: str | None = None, by: str = "Dana Scully"):
    return {
        "id": item_id,
        "name": name,
        "size": 2048,
        "webUrl": f"{SITE_URL}/Shared%20Documents/Temp%20Folder%201/{name.replace(' ', '%20')}",
        "file": {"mimeType": "application/octet-stream"},
        "createdDateTime": created or modified,
        "lastModifiedDateTime": modified,
        "createdBy": {"user": {"displayName": by}},
        "lastModifiedBy": {"user": {"displayName": by}},
    }


def graph_folder(item_id: str, name: str) -> dict:
    return {
        "id": item_id,
        "name": name,
        "folder": {"childCount": 1},
        "webUrl": f"{SITE_URL}/Shared%20Documents/Temp%20Folder%201/{name}",
        "createdDateTime": "2026-01-01T00:00:00Z",
        "lastModifiedDateTime": "2026-01-01T00:00:00Z",
    }


class FakeModel:
    """Records prompts; answers digest prompts with a digest and everything else with a numbered summary."""

    deployment = "gpt-test"

    def __init__(self):
        self.prompts: list[str] = []
        self._lock = threading.Lock()

    def complete(self, system: str, prompt: str) -> str:
        with self._lock:
            self.prompts.append(prompt)
            number = len(self.prompts)
        if "Write a digest" in prompt:
            return "## Highlights\n- Launch is on track (Q3 plan.pptx)"
        return f"**Overview:** Summary number {number}."


def make_settings(**overrides) -> Settings:
    return Settings(
        sharepoint=SharePointSettings(
            site_url=SITE_URL,
            library="Documents",
            auth_mode="app",
            tenant_id="tenant",
            client_id="client",
            client_secret="secret",
        ),
        foundry=FoundrySettings(
            endpoint=overrides.get("endpoint", "my-resource"),
            deployment=overrides.get("deployment", "gpt-4.1"),
            api_key=overrides.get("api_key", "test-key"),
            reasoning_effort=overrides.get("reasoning_effort"),
            max_output_tokens=overrides.get("max_output_tokens"),
        ),
        max_input_chars=overrides.get("max_input_chars", 200_000),
    )
