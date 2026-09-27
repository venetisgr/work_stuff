"""Turn PowerPoint and Word files into plain text the model can read."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.shapes.group import GroupShape

_HEADING_STYLE = re.compile(r"Heading (\d)")
_MAX_CHART_POINTS = 24  # per series; larger charts are listed by series name only


class ExtractionError(Exception):
    """The file couldn't be opened: password-protected, encrypted by a sensitivity label, or damaged."""


def extract_text(data: bytes, extension: str) -> str:
    readers = {".pptx": pptx_text, ".pptm": pptx_text, ".docx": docx_text}
    if extension not in readers:
        raise ValueError(f"Unsupported file type: {extension}")
    try:
        return readers[extension](data)
    except ExtractionError:
        raise
    except Exception as exc:  # a layout the parser doesn't handle
        raise ExtractionError(f"Couldn't read the file's content ({type(exc).__name__}: {exc}).") from exc


def _open(opener, data: bytes):
    try:
        return opener(io.BytesIO(data))
    except Exception as exc:  # zipfile, lxml and the parsers raise a variety of errors on unreadable files
        raise ExtractionError(
            "Couldn't open the file; it may be password-protected, encrypted by a sensitivity label, or damaged."
        ) from exc


def pptx_text(data: bytes) -> str:
    """One "## Slide N: title" section per slide: text, tables, chart data and speaker notes."""
    presentation = _open(Presentation, data)
    sections = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None else ""
        lines = [
            line
            for shape in slide.shapes
            if title_shape is None or shape.shape_id != title_shape.shape_id
            for line in _shape_lines(shape)
        ]
        notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = _clean(notes_frame.text) if notes_frame is not None else ""
        if not (title or lines or notes):
            continue
        section = [f"## Slide {number}" + (f": {title}" if title else ""), *lines]
        if notes:
            section.append(f"Speaker notes: {notes}")
        sections.append("\n".join(section))
    return "\n\n".join(sections)


def _shape_lines(shape) -> list[str]:
    if isinstance(shape, GroupShape):
        return [line for child in shape.shapes for line in _shape_lines(child)]
    if shape.has_text_frame:
        return [
            "  " * paragraph.level + "- " + text
            for paragraph in shape.text_frame.paragraphs
            if (text := _clean(paragraph.text))
        ]
    if shape.has_table:
        return _rows_to_lines([_clean(cell.text) for cell in row.cells] for row in shape.table.rows)
    if shape.has_chart:
        return _chart_lines(shape.chart)
    return []


def _chart_lines(chart) -> list[str]:
    try:
        title = ""
        if chart.has_title and chart.chart_title.has_text_frame:
            title = _clean(chart.chart_title.text_frame.text)
        lines = [f"[Chart: {title}]" if title else "[Chart]"]
        for plot in chart.plots:
            categories = [str(category) for category in plot.categories]
            for series in plot.series:
                values = list(series.values)
                if categories and len(values) == len(categories) <= _MAX_CHART_POINTS:
                    points = "; ".join(f"{c} = {_number(v)}" for c, v in zip(categories, values, strict=True))
                    lines.append(f"  {series.name or 'Series'}: {points}")
                elif series.name:
                    lines.append(f"  Series: {series.name}")
        return lines
    except Exception:  # python-pptx doesn't understand every chart type
        return ["[Chart]"]


def _number(value: float | None) -> str:
    if value is None:
        return "n/a"
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def docx_text(data: bytes) -> str:
    """Paragraphs and tables in document order; headings become Markdown headings, list items bullets."""
    document = _open(Document, data)
    lines = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if text := _clean(block.text):
                lines.append(_paragraph_line(block, text))
        else:
            lines.extend(_docx_table_lines(block))
    return "\n".join(lines)


def _paragraph_line(paragraph: Paragraph, text: str) -> str:
    style = (paragraph.style.name if paragraph.style is not None else None) or ""
    if style == "Title":
        return f"# {text}"
    if match := _HEADING_STYLE.fullmatch(style):
        return "#" * min(int(match.group(1)) + 1, 6) + " " + text
    properties = paragraph._p.pPr
    if style.startswith("List") or (properties is not None and properties.numPr is not None):
        return f"- {text}"
    return text


def _docx_table_lines(table: Table) -> list[str]:
    try:
        rows = []
        for row in table.rows:
            cells, previous = [], None
            for cell in row.cells:
                if previous is not None and cell._tc is previous._tc:  # merged cells repeat
                    continue
                previous = cell
                cells.append(_clean(cell.text))
            rows.append(cells)
        return _rows_to_lines(rows)
    except Exception:  # unusual table layouts can trip python-docx; keep the rest of the document
        return ["[Table that couldn't be read]"]


def _rows_to_lines(rows) -> list[str]:
    return [" | ".join(cells) for cells in rows if any(cells)]


def _clean(text: str) -> str:
    return " ".join(text.split())
